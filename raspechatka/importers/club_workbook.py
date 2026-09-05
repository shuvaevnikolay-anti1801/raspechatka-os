"""One-time local import from the legacy Club Google Sheets XLSX export.

Run only inside bench, after copying the XLSX file into the container:
bench --site <site> execute raspechatka.importers.club_workbook.execute --kwargs '{"path":"/path/file.xlsx","default_business_point":"POINT-CODE","limit":3}'
"""

import frappe
from frappe.utils import cint, get_datetime
from openpyxl import load_workbook

from raspechatka.raspechatka_os.doctype.client.client import normalize_phone


def execute(path, default_business_point, dry_run=False, limit=None):
	point = default_business_point if frappe.db.exists("Business Point", default_business_point) else frappe.db.get_value("Business Point", {"point_code": default_business_point}, "name")
	if not point:
		raise ValueError("Не найдена точка продаж для старых клиентов")
	with open(path, "rb") as source:
		workbook = load_workbook(source, data_only=True, read_only=False)
	clients = _rows(workbook["Клиенты"])
	if cint(limit) > 0:
		clients = clients[:cint(limit)]
	channels = _rows(workbook["Каналы"])
	consents = _rows(workbook["Согласия"])
	channels_by_client = _latest_channels(channels)
	consents_by_client = {}
	for row in consents:
		if row.get("Consent ID") and row.get("Client ID"):
			consents_by_client.setdefault(str(row["Client ID"]), []).append(row)
	result = {"created": 0, "updated": 0, "skipped": 0, "consents": 0}
	for row in clients:
		legacy_id = str(row.get("Client ID") or "").strip()
		phone = normalize_phone(row.get("Телефон"))
		if not legacy_id or not phone:
			result["skipped"] += 1
			continue
		name = frappe.db.get_value("Client", {"phone": phone}, "name")
		doc = frappe.get_doc("Client", name) if name else frappe.new_doc("Client")
		doc.client_id = doc.client_id or legacy_id
		doc.phone = phone
		doc.first_name = str(row.get("Имя") or "Клиент").strip()
		doc.registration_point = doc.registration_point or point
		doc.registration_source = row.get("Источник") or "Google Sheets — Клуб Распечатка"
		doc.registered_at = row.get("Дата регистрации") or doc.registered_at
		doc.personal_data_consent = cint(row.get("Согласие ПД"))
		doc.marketing_consent = cint(row.get("Согласие на рекламу"))
		doc.club_rules_consent = cint(row.get("Правила клуба"))
		doc.link_token = row.get("Link Token") or doc.link_token
		doc.notes = row.get("Комментарий") or doc.notes
		doc.set("messengers", [])
		for channel in channels_by_client.get(legacy_id, {}).values():
			doc.append("messengers", {
				"messenger_type": channel.get("Канал"), "platform_user_id": _as_id(channel.get("Platform User ID")),
				"bothelp_subscriber_id": _as_id(channel.get("BotHelp Subscriber ID")), "status": channel.get("Статус") or "Отключен",
				"connected_at": channel.get("Дата подключения"), "last_activity": channel.get("Последняя активность"),
				"health": channel.get("Health") or "ACTIVE", "last_error": channel.get("Последняя ошибка"),
				"disconnected_at": channel.get("Дата отключения"),
			})
		if not dry_run:
			doc.save(ignore_permissions=True)
		result["updated" if name else "created"] += 1
		if dry_run:
			continue
		for consent in consents_by_client.get(legacy_id, []):
			if frappe.db.exists("Client Consent", {"client": doc.name, "submission_id": consent.get("Form Submission ID"), "consent_type": consent.get("Тип согласия")}):
				continue
			frappe.get_doc({
				"doctype": "Client Consent", "client": doc.name, "consent_type": consent.get("Тип согласия"),
				"accepted": cint(consent.get("Согласие")), "recorded_at": consent.get("Дата/время"),
				"document_version": consent.get("Версия документа") or "legacy", "document_url": consent.get("URL документа"),
				"source": consent.get("Источник") or "Google Sheets", "ip_address": consent.get("IP"),
				"user_agent": consent.get("User-Agent"), "submission_id": consent.get("Form Submission ID"),
				"revoked_at": consent.get("Дата отзыва"),
			}).insert(ignore_permissions=True)
			result["consents"] += 1
	if not dry_run:
		frappe.db.commit()
	return result


def _rows(sheet):
	values = list(sheet.iter_rows(values_only=True))
	if not values:
		return []
	headers = list(values[0])
	return [{headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]} for row in values[1:] if any(value not in (None, "") for value in row)]


def _latest_channels(rows):
	result = {}
	for row in rows:
		client_id = str(row.get("Client ID") or "").strip()
		channel = str(row.get("Канал") or "").strip()
		if not client_id or channel not in ("Telegram", "MAX", "VK"):
			continue
		current = result.setdefault(client_id, {}).get(channel)
		current_date = current.get("Последняя активность") if current else None
		row_date = row.get("Последняя активность")
		if not current or (row_date and (not current_date or get_datetime(row_date) >= get_datetime(current_date))):
			result[client_id][channel] = row
	return result


def _as_id(value):
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value or "")
